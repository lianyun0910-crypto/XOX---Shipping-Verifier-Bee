from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

from .normalizer import FIELDS
from .storage import RECORD_DIR, save_excel

EXCEL_PATH = RECORD_DIR / "shipping_verification_audit.xlsx"

SHEET_HEADERS = {
    "Summary": ["run_id", "timestamp_utc", "source", "total", "ok", "mismatch", "needs_review", "documents_processed", "fields_compared", "cloud_status", "cloud_file"],
    "Email Classification": ["run_id", "timestamp_utc", "filename", "email_id", "from", "subject", "language", "category", "confidence", "reason", "explanation", "evidence", "next_action", "status"],
    "Document Classification": ["run_id", "timestamp_utc", "email_id", "filename", "document_type", "language", "confidence", "shipment_reference", "booking_reference", "status", "review_reason", "explanation", "evidence"],
    "Field Comparison": ["run_id", "timestamp_utc", "email_id", "source_filename", "si_filename", "bl_filename", "overall_status", "overall_confidence", "field", "si_value", "bl_value", "field_status", "reason", "evidence"],
    "Audit Log": ["run_id", "timestamp_utc", "source", "email_id", "filename", "category", "document_type", "ai_confidence", "status", "final_status", "review_reason", "human_decision", "reviewer", "review_note", "details"],
}

HEADER_FILL = PatternFill("solid", fgColor="172033")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
OK_FILL = PatternFill("solid", fgColor="E8F7F1")
OK_FONT = Font(color="07835F", bold=True)
MISMATCH_FILL = PatternFill("solid", fgColor="FFF0F0")
MISMATCH_FONT = Font(color="BB4B4B", bold=True)
REVIEW_FILL = PatternFill("solid", fgColor="FFF6E6")
REVIEW_FONT = Font(color="A96E1E", bold=True)
LINK_FONT = Font(color="0563C1", underline="single", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="E3E8EF"))


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def _status(value):
    return str(value or "").strip().upper()


def _doc_filename(doc):
    return str(doc.get("filename") or doc.get("name") or "")


def _doc_result(doc):
    return doc.get("result") or {}


def _doc_type(doc):
    result = _doc_result(doc)
    return str(doc.get("document_type") or result.get("document_type") or "")


def _doc_language(doc):
    return str(_doc_result(doc).get("language") or "")


def _shipment(doc):
    return str(_doc_result(doc).get("shipment_reference") or "")


def _booking(doc):
    return str(_doc_result(doc).get("booking_reference") or "")


def _ensure_workbook():
    RECORD_DIR.mkdir(parents=True, exist_ok=True)
    if EXCEL_PATH.exists():
        wb = load_workbook(EXCEL_PATH)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for name, headers in SHEET_HEADERS.items():
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            ws.append(headers)
        else:
            ws = wb[name]
            existing = [c.value for c in ws[1]]
            for header in headers:
                if header not in existing:
                    ws.cell(1, ws.max_column + 1, header)
    return wb


def _style(ws, status_column=None):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 28
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BORDER
    ws.freeze_panes = "A2"
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:{ws.cell(1, ws.max_column).column_letter}{ws.max_row}"
        if ws.tables:
            ws.tables.clear()
        table = Table(displayName=f"Table_{ws.title.replace(' ', '_').replace('/', '_')}", ref=f"A1:{ws.cell(1, ws.max_column).column_letter}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    if status_column:
        headers = [c.value for c in ws[1]]
        if status_column in headers:
            col = headers.index(status_column) + 1
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col)
                value = _status(cell.value)
                if value in {"OK", "MATCH"}:
                    cell.fill = OK_FILL
                    cell.font = OK_FONT
                elif value == "MISMATCH":
                    cell.fill = MISMATCH_FILL
                    cell.font = MISMATCH_FONT
                elif value in {"NEEDS_REVIEW", "UNRESOLVED", "UNKNOWN"}:
                    cell.fill = REVIEW_FILL
                    cell.font = REVIEW_FONT
    if ws.title == "Summary":
        headers = [c.value for c in ws[1]]
        if "cloud_file" in headers:
            col = headers.index("cloud_file") + 1
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col)
                if cell.hyperlink:
                    cell.font = LINK_FONT


def _set_widths(ws):
    maps = {
        "Summary": [34, 25, 14, 10, 10, 12, 14, 19, 16, 16, 24],
        "Email Classification": [34, 25, 32, 28, 30, 60, 14, 24, 13, 50, 90, 70, 60, 18],
        "Document Classification": [34, 25, 28, 38, 28, 14, 13, 25, 25, 20, 30, 80, 70],
        "Field Comparison": [34, 25, 28, 35, 35, 35, 18, 16, 25, 55, 55, 18, 80, 70],
        "Audit Log": [34, 25, 14, 28, 36, 24, 28, 16, 18, 18, 28, 24, 24, 50, 80],
    }
    widths = maps.get(ws.title, [])
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = width


def _counts(results):
    documents_processed = 0
    fields_compared = 0
    for result in results:
        seen = set()
        for doc in result.get("documents") or []:
            name = _doc_filename(doc)
            key = (name.lower(), _doc_type(doc).upper())
            if name and key not in seen:
                seen.add(key)
                documents_processed += 1
        comparison = result.get("comparison") or {}
        fields_compared += len(comparison.get("fields") or {})
    return {
        "total": len(results),
        "ok": sum(_status(r.get("status")) == "OK" for r in results),
        "mismatch": sum(_status(r.get("status")) == "MISMATCH" for r in results),
        "needs_review": sum(_status(r.get("status")) == "NEEDS_REVIEW" for r in results),
        "documents_processed": documents_processed,
        "fields_compared": fields_compared,
    }


def append_run(run_id: str, source: str, results: list[dict]) -> dict:
    wb = _ensure_workbook()
    timestamp = datetime.now(timezone.utc).isoformat()
    counts = _counts(results)
    ws = wb["Summary"]
    ws.append([run_id, timestamp, source, counts["total"], counts["ok"], counts["mismatch"], counts["needs_review"], counts["documents_processed"], counts["fields_compared"], "Uploading...", "Preparing Cloud File..."])
    summary_row = ws.max_row

    for result in results:
        filename = str(result.get("filename", ""))
        email_id = str(result.get("email_id", ""))
        category = str(result.get("category", result.get("email_category", "")))
        overall_status = str(result.get("status", ""))
        review_reason = str(result.get("review_reason", ""))
        reason = str(result.get("reason", ""))
        details = reason or str(result.get("details", ""))
        human = result.get("human_review") or {}
        wb["Audit Log"].append([
            run_id, timestamp, source, email_id, filename, category,
            result.get("document_type", ""), result.get("decision_confidence", result.get("confidence", "")),
            overall_status, result.get("final_status", ""), review_reason,
            human.get("decision", ""), human.get("reviewer", ""), human.get("review_note", ""),
            details,
        ])
        if result.get("email_category"):
            evidence = result.get("evidence") or []
            wb["Email Classification"].append([
                run_id, timestamp, filename, email_id, result.get("from", ""), result.get("subject", ""),
                result.get("language", ""), result.get("email_category", ""), result.get("classification_confidence", result.get("confidence", "")),
                result.get("reason", ""), result.get("explanation", ""), ", ".join(map(str, evidence)),
                result.get("next_action", ""), overall_status,
            ])

        documents = list(result.get("documents") or [])
        if result.get("si"):
            documents.append(result["si"])
        if result.get("bl"):
            documents.append(result["bl"])
        seen_documents = set()
        for doc in documents:
            name = _doc_filename(doc)
            dtype = _doc_type(doc)
            key = (name.lower(), dtype.upper())
            if not name or key in seen_documents:
                continue
            seen_documents.add(key)
            doc_result = _doc_result(doc)
            wb["Document Classification"].append([
                run_id, timestamp, email_id, name, dtype, _doc_language(doc), doc_result.get("confidence", ""),
                _shipment(doc), _booking(doc), doc.get("status", ""), doc.get("review_reason", ""),
                doc_result.get("explanation", ""), ", ".join(map(str, doc_result.get("evidence") or [])),
            ])

        comparison = result.get("comparison") or {}
        fields = comparison.get("fields") or {}
        if fields:
            si_filename = str(result.get("si_filename", ""))
            bl_filename = str(result.get("bl_filename", ""))
            for field in FIELDS:
                item = fields.get(field) or {}
                wb["Field Comparison"].append([
                    run_id, timestamp, email_id, filename, si_filename, bl_filename,
                    comparison.get("status", overall_status), comparison.get("confidence", ""), field,
                    _clean(item.get("si")), _clean(item.get("bl")),
                    item.get("status", "UNRESOLVED"), item.get("reason", ""), item.get("evidence", ""),
                ])

    for name in SHEET_HEADERS:
        _style(wb[name], {"Email Classification": "status", "Document Classification": "status", "Field Comparison": "field_status", "Audit Log": "status"}.get(name))
        _set_widths(wb[name])

    wb.properties.title = "Shipping Document Verification Audit"
    wb.properties.creator = "Shipping Verifier"
    wb.save(EXCEL_PATH)

    cloud = save_excel(EXCEL_PATH)
    uploaded = bool(cloud.get("uploaded"))
    cloud_url = cloud.get("cloud_url")

    cloud_status_cell = ws.cell(summary_row, 10)
    cloud_file_cell = ws.cell(summary_row, 11)
    if uploaded:
        cloud_status_cell.value = "Uploaded"
        cloud_status_cell.fill = OK_FILL
        cloud_status_cell.font = OK_FONT
        if cloud_url:
            cloud_file_cell.value = "Open Cloud File"
            cloud_file_cell.hyperlink = cloud_url
            cloud_file_cell.font = LINK_FONT
    else:
        cloud_status_cell.value = "Local Only"
        cloud_status_cell.fill = REVIEW_FILL
        cloud_status_cell.font = REVIEW_FONT
        cloud_file_cell.value = ""
        cloud_file_cell.hyperlink = None

    wb.save(EXCEL_PATH)
    try:
        final_cloud = save_excel(EXCEL_PATH)
        final_cloud_url = final_cloud.get("cloud_url") or cloud_url
    except Exception:
        final_cloud_url = cloud_url

    return {
        "local_path": str(EXCEL_PATH),
        "filename": EXCEL_PATH.name,
        "cloud_url": final_cloud_url,
        "uploaded": uploaded,
        "cloud_status": "Uploaded" if uploaded else "Local Only",
    }
