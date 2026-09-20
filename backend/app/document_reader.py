import csv
import io
import os
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")

def read_csv(path: Path) -> str:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    rows = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in rows)

def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append(text)
    return "\n".join(pages)

def read_docx(path: Path) -> str:
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)

def read_xlsx(path: Path) -> str:
    wb = load_workbook(filename=path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                out.append(" | ".join(vals))
    return "\n".join(out)

def read_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                out.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)

def read_image(path: Path) -> str:
    if os.getenv("OCR_ENABLED", "false").lower() != "true":
        return ""
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(path))
    except Exception:
        return ""

def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".log"}:
        return read_txt(path)
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    if suffix == ".pptx":
        return read_pptx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif"}:
        return read_image(path)
    raise ValueError(f"Unsupported document format: {suffix}")
